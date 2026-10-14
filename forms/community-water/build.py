"""Rebuild the Bo town pilot water-service XLSForm (requires openpyxl)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
ROOT = Path(__file__).resolve().parent
cols = ['type','name','label','hint','required','relevant','constraint','constraint_message','default','appearance','parameters','read_only']
rows = []
def add(t,n,l='',**kw): rows.append(dict(type=t,name=n,label=l,**kw))
add('start','start'); add('end','end')
add('note','intro','Bo town water and sanitation pilot. This survey covers Bo town only, not the rest of Bo District. Explain the purpose and voluntary participation. Do not collect names, phone numbers or personal documents.')
add('select_one yes_no','consent','Has the respondent agreed to participate?',required='yes')
add('note','declined','Thank the respondent and end the interview.',relevant="${consent} = 'no'")
add('select_one yes_no','in_bo_town','Is this interview site within Bo town?',required='yes',relevant="${consent} = 'yes'",hint='Confirm with the respondent or local supervisor. Do not include other towns or villages in Bo District.')
add('note','outside_scope','This site is outside the Bo town survey area. End this interview; do not collect water-service details.',relevant="${consent} = 'yes' and ${in_bo_town} = 'no'")
add('begin_group','visit','Bo town water service',relevant="${consent} = 'yes' and ${in_bo_town} = 'yes'")
add('select_one interview_unit','interview_unit','Reporting unit',required='yes',hint='Keep household reports separate from community-representative reports in analysis.')
add('date','visit_date','Visit date',required='yes',default='today()')
add('select_one district','district','District',required='yes',default='bo',read_only='yes')
add('text','town','Town',default='Bo',read_only='yes')
add('text','community','Neighbourhood / section in Bo town',hint='Use the locally recognized name. Ask the supervisor if uncertain; do not guess a ward number.',required='yes',constraint='string-length(.) <= 150',constraint_message='Use at most 150 characters.')
add('text','site_code','Water point or site reference',required='yes',hint='Use your project reference, for example BO-WASH-0001. Do not use a person’s name.',constraint="regex(., '^[A-Za-z0-9_-]{3,40}$')",constraint_message='Use 3–40 letters, numbers, underscores or hyphens.')
add('select_one water_source','water_source','Main drinking-water source',required='yes')
add('text','other_source','Specify the other source',required='yes',relevant="${water_source} = 'other'")
add('select_one condition','functioning','Is the main water source functioning today?',required='yes')
add('integer','round_trip_minutes','Reported round-trip time to collect water (minutes)',hint='Include waiting time. Leave blank if unknown.',constraint='. >= 0 and . <= 1440',constraint_message='Enter 0–1,440 minutes.')
add('select_one yes_no_unknown','year_round','Is water reportedly available throughout the year?',required='yes')
add('select_one yes_no_unknown','safe_storage','Does the respondent report using a covered water-storage container?',required='yes')
add('select_one sanitation','sanitation','Main sanitation facility used by the community/household interviewed',required='yes')
add('select_one yes_no_unknown','handwashing','Is a handwashing place with water and soap available?',required='yes')
add('select_multiple needs','priority_needs','Reported priority service needs',required='yes',constraint="not(selected(., 'none') and count-selected(.) > 1)",constraint_message='Select None on its own.')
add('text','other_need','Specify another service need',required='yes',relevant="selected(${priority_needs}, 'other')")
add('select_one yes_no','location_consent','May the water-point location be recorded?',required='yes',hint='Explain that authorized project reviewers can view this location. Refusal does not end the interview.')
add('geopoint','water_point_location','Optional water-point location',relevant="${location_consent} = 'yes'",hint='Record the water point, not a private home. Location accuracy is device-reported.')
add('image','water_point_photo','Optional water-point photo',hint='Obtain permission. Avoid people and personal documents. Leave blank if inappropriate.',parameters='max-pixels=1600')
add('text','notes','Observations and limitations',constraint='string-length(.) <= 1000',constraint_message='Use at most 1,000 characters.',appearance='multiline')
add('note','interpretation','This pilot records observations and respondent reports. It does not test water quality or certify safety. Follow your organization’s referral procedure for reported service problems.')
add('end_group','visit')
choices=[]
def choice(listname,name,label): choices.append([listname,name,label])
for n,l in [('yes','Yes'),('no','No')]: choice('yes_no',n,l); choice('yes_no_unknown',n,l)
choice('yes_no_unknown','unknown','Unknown')
choice('interview_unit','household','Household respondent'); choice('interview_unit','community','Community representative')
choice('district','bo','Bo')
for n,l in [('piped','Piped supply'),('borehole','Borehole / tube well'),('protected_well','Protected well'),('unprotected_well','Unprotected well'),('spring','Spring'),('rainwater','Rainwater'),('surface','River, stream, pond or other surface water'),('vendor','Water vendor / tanker'),('packaged','Bottled or sachet water'),('other','Other')]: choice('water_source',n,l)
for n,l in [('yes','Functioning'),('partial','Partly functioning / intermittent'),('no','Not functioning'),('unknown','Unknown')]: choice('condition',n,l)
for n,l in [('flush','Flush or pour-flush toilet'),('pit_slab','Pit latrine with slab'),('pit_no_slab','Pit latrine without slab'),('shared','Shared / communal facility'),('none','No facility'),('other','Other / unknown')]: choice('sanitation',n,l)
for n,l in [('repair','Repair the water point'),('reliability','Improve reliability / seasonal access'),('distance','Reduce collection distance or waiting time'),('storage','Improve storage'),('sanitation','Improve sanitation'),('handwashing','Improve handwashing facilities'),('other','Other'),('none','None reported')]: choice('needs',n,l)
book=Workbook(); book.remove(book.active)
for name, headers, data in [('survey',cols,[[r.get(c,'') for c in cols] for r in rows]),('choices',['list_name','name','label'],choices),('settings',['form_title','form_id','version','instance_name'],[['Bo Town: Water and Sanitation','sl_community_water_service','2026101402',"if(${consent} = 'yes', if(${in_bo_town} = 'yes', concat(${site_code}, ' - ', ${community}, ' - Bo'), 'Outside Bo town'), 'Participation declined')"]])]:
 sheet=book.create_sheet(name);sheet.append(headers)
 for r in data:sheet.append(r)
 for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='0D748B')
 sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
 for column in sheet.columns:
  letter=column[0].column_letter;sheet.column_dimensions[letter].width=min(65,max(18,max(len(str(c.value or '')) for c in column)+2))
 for row_cells in sheet.iter_rows(min_row=2):
  for c in row_cells:c.alignment=Alignment(wrap_text=True,vertical='top')
book.save(ROOT/'sl_community_water_service.xlsx')
