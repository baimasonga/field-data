"""Build the Field Data test survey XLSForms.

Run from this folder: python3 build.py
Writes agriculture, census, census back-check and health XLSForms next to
this script. The forms are test material: realistic, but not validated
instruments.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
VERSION = '2026101101'

SURVEY_COLS = ['type', 'name', 'label', 'hint', 'required', 'relevant',
               'constraint', 'constraint_message', 'calculation', 'default',
               'appearance', 'parameters', 'repeat_count', 'choice_filter', 'read_only']

DISTRICTS = [
    ('bo', 'Bo'), ('bombali', 'Bombali'), ('bonthe', 'Bonthe'), ('falaba', 'Falaba'),
    ('kailahun', 'Kailahun'), ('kambia', 'Kambia'), ('karene', 'Karene'),
    ('kenema', 'Kenema'), ('koinadugu', 'Koinadugu'), ('kono', 'Kono'),
    ('moyamba', 'Moyamba'), ('port_loko', 'Port Loko'), ('pujehun', 'Pujehun'),
    ('tonkolili', 'Tonkolili'), ('western_rural', 'Western Area Rural'),
    ('western_urban', 'Western Area Urban'),
]
YES_NO = [('yes', 'Yes'), ('no', 'No')]
YES_NO_DK = YES_NO + [('dont_know', "Don't know")]
SEX = [('male', 'Male'), ('female', 'Female')]
PHONE = "regex(., '^0[0-9]{8}$')"
PHONE_MSG = 'Enter 9 digits starting with 0, for example 076123456.'


def q(type_, name, label='', **kw):
    row = {'type': type_, 'name': name, 'label': label}
    row.update(kw)
    return row


def start_meta():
    return [q('start', 'start'), q('end', 'end'), q('today', 'today'),
            q('deviceid', 'deviceid'), q('username', 'username')]


def consent(intro):
    return [
        q('note', 'intro', intro),
        q('select_one yes_no', 'consent',
          'Has the respondent agreed to take part in this interview?', required='yes'),
        q('note', 'no_consent', 'Thank the respondent and end the interview. Save the form as complete.',
          relevant="${consent} = 'no'"),
    ]


def location(code_name, code_label, code_regex, code_msg, code_hint):
    return [
        q('begin_group', 'location', 'Location'),
        q('select_one district', 'district', 'District', required='yes'),
        q('text', 'chiefdom', 'Chiefdom', required='yes'),
        q('text', 'village', 'Town or village', required='yes'),
        q('text', code_name, code_label, hint=code_hint, required='yes',
          constraint=f"regex(., '{code_regex}')", constraint_message=code_msg),
        q('geopoint', 'gps', 'Record the GPS location',
          hint='Stand at the place itself and wait until the accuracy is 10 m or better.',
          required='yes'),
        q('end_group', 'location'),
    ]


# ---------------------------------------------------------------- agriculture
CROPS = [
    ('cocoa', 'Cocoa'), ('coffee', 'Coffee'), ('oil_palm', 'Oil palm'), ('cashew', 'Cashew'),
    ('rice', 'Rice'), ('cassava', 'Cassava'), ('maize', 'Maize'), ('groundnut', 'Groundnut'),
    ('vegetables', 'Vegetables'), ('other', 'Other'),
]
agriculture = {
    'file': 'fd_test_agriculture.xlsx',
    'settings': {'form_title': 'Test: Farm household and plot survey',
                 'form_id': 'fd_test_agriculture',
                 'instance_name': "concat(${farm_id}, ' - ', ${farmer_name})"},
    'survey': start_meta() + consent(
        'Field Data TEST form. Farm household and plot survey. '
        'Introduce yourself and explain that the interview takes about 20 minutes.') + [
        q('begin_group', 'main', '', relevant="${consent} = 'yes'"),
        q('text', 'enumerator', 'Enumerator name', required='yes'),
        *location('farm_id', 'Farm ID', '^FARM-[0-9]{4}$',
                  'Use the format FARM-0001.', 'Format FARM-0001, from the farm register.'),
        q('begin_group', 'farmer', 'Farmer'),
        q('text', 'farmer_name', 'Full name of the farmer', required='yes'),
        q('select_one sex', 'farmer_sex', 'Sex of the farmer', required='yes'),
        q('integer', 'farmer_age', 'Age of the farmer (years)', required='yes',
          constraint='. >= 15 and . <= 100', constraint_message='Age must be between 15 and 100.'),
        q('text', 'farmer_phone', 'Phone number', hint='Leave blank if none.',
          constraint=PHONE, constraint_message=PHONE_MSG, appearance='numbers'),
        q('integer', 'hh_size', 'How many people live in the household?', required='yes',
          constraint='. >= 1 and . <= 40', constraint_message='Between 1 and 40.'),
        q('select_one yes_no', 'coop_member', 'Is the farmer a member of a cooperative or farmer group?',
          required='yes'),
        q('text', 'coop_name', 'Name of the cooperative or group', relevant="${coop_member} = 'yes'"),
        q('end_group', 'farmer'),

        q('begin_group', 'crops_grp', 'Crops'),
        q('select_multiple crop', 'main_crops', 'Which crops did the household grow in the last 12 months?',
          required='yes'),
        q('text', 'other_crop', 'Other crop: specify', required='yes',
          relevant="selected(${main_crops}, 'other')"),
        q('integer', 'plot_count', 'How many plots does the household farm?', required='yes',
          constraint='. >= 1 and . <= 10', constraint_message='Between 1 and 10.'),
        q('end_group', 'crops_grp'),

        q('begin_repeat', 'plot', 'Plot', repeat_count='${plot_count}'),
        q('calculate', 'plot_no', calculation='position(..)'),
        q('note', 'plot_title', 'Plot ${plot_no}'),
        q('select_one crop', 'plot_crop', 'Main crop on this plot', required='yes',
          choice_filter="selected(${main_crops}, name)"),
        q('decimal', 'plot_area_ha', 'Plot area (hectares)', required='yes',
          constraint='. > 0 and . <= 100', constraint_message='Between 0.01 and 100 hectares.'),
        q('select_one tenure', 'plot_tenure', 'How does the household hold this plot?', required='yes'),
        q('select_one yes_no', 'plot_irrigated', 'Is the plot irrigated?', required='yes'),
        q('geoshape', 'plot_boundary', 'Walk the plot boundary (optional)'),
        q('end_repeat', 'plot'),
        q('calculate', 'total_area_ha', calculation='sum(${plot_area_ha})'),
        q('note', 'total_area_note', 'Total area farmed: ${total_area_ha} ha'),

        q('begin_group', 'harvest', 'Last season'),
        q('decimal', 'harvest_kg', 'Total harvest of the main crop last season (kg)', required='yes',
          constraint='. >= 0 and . <= 200000', constraint_message='Between 0 and 200,000 kg.'),
        q('decimal', 'sold_kg', 'How much of it was sold (kg)?', required='yes',
          constraint='. >= 0 and . <= ${harvest_kg}',
          constraint_message='Cannot sell more than was harvested.'),
        q('integer', 'price_le', 'Average price received per kg (NLe)',
          relevant='${sold_kg} > 0', required='yes',
          constraint='. > 0 and . <= 1000', constraint_message='Between 1 and 1,000 NLe per kg.'),
        q('calculate', 'sales_le', calculation='if(${sold_kg} > 0, ${sold_kg} * ${price_le}, 0)'),
        q('note', 'sales_note', 'Estimated sales: NLe ${sales_le}', relevant='${sold_kg} > 0'),
        q('end_group', 'harvest'),

        q('begin_group', 'inputs', 'Inputs and support'),
        q('select_one yes_no', 'used_fertilizer', 'Did the household use fertilizer last season?',
          required='yes'),
        q('select_multiple fertilizer', 'fertilizer_type', 'Which fertilizer?', required='yes',
          relevant="${used_fertilizer} = 'yes'"),
        q('select_multiple pest', 'pests', 'Which pests or diseases affected the crops?', required='yes',
          constraint="not(selected(., 'none')) or count-selected(.) = 1",
          constraint_message='"None" cannot be chosen with other answers.'),
        q('select_one yes_no', 'extension_visit',
          'Did an extension worker visit in the last 12 months?', required='yes'),
        q('integer', 'extension_visits', 'How many visits?', required='yes',
          relevant="${extension_visit} = 'yes'",
          constraint='. >= 1 and . <= 52', constraint_message='Between 1 and 52.'),
        q('image', 'farm_photo', 'Photo of the main plot (optional)', parameters='max-pixels=1024'),
        q('end_group', 'inputs'),
        q('text', 'remarks', 'Enumerator remarks (optional)', appearance='multiline'),
        q('end_group', 'main'),
    ],
    'choices': {
        'yes_no': YES_NO, 'sex': SEX, 'district': DISTRICTS, 'crop': CROPS,
        'tenure': [('owned', 'Owned'), ('family', 'Family land'), ('rented', 'Rented'),
                   ('sharecropped', 'Sharecropped'), ('other', 'Other')],
        'fertilizer': [('npk', 'NPK'), ('urea', 'Urea'), ('organic', 'Organic or compost'),
                       ('other', 'Other')],
        'pest': [('none', 'None'), ('fall_armyworm', 'Fall armyworm'),
                 ('black_pod', 'Black pod (cocoa)'), ('capsids', 'Capsids'),
                 ('rice_blast', 'Rice blast'), ('rodents', 'Rodents'), ('other', 'Other')],
    },
}

# --------------------------------------------------------------------- census
RELATIONSHIP = [
    ('head', 'Head'), ('spouse', 'Spouse'), ('child', 'Son or daughter'),
    ('grandchild', 'Grandchild'), ('parent', 'Parent'), ('sibling', 'Brother or sister'),
    ('other_relative', 'Other relative'), ('not_related', 'Not related'),
]
census = {
    'file': 'fd_test_census.xlsx',
    'settings': {'form_title': 'Test: Household census',
                 'form_id': 'fd_test_census',
                 'instance_name': "concat(${household_id}, ' - ', ${respondent_name})"},
    'survey': start_meta() + consent(
        'Field Data TEST form. Household census. Ask to speak to the head of household '
        'or another adult member who knows everyone in the household.') + [
        q('begin_group', 'main', '', relevant="${consent} = 'yes'"),
        q('text', 'enumerator', 'Enumerator name', required='yes'),
        *location('household_id', 'Household ID', '^HH-[0-9]{6}$',
                  'Use the format HH-000001.', 'Format HH-000001, from the listing sheet.'),
        q('text', 'ea_code', 'Enumeration area code', required='yes',
          constraint="regex(., '^[0-9]{4}$')", constraint_message='4 digits.'),

        q('begin_group', 'respondent', 'Respondent'),
        q('text', 'respondent_name', 'Respondent name', required='yes'),
        q('text', 'respondent_phone', 'Phone number', hint='Leave blank if none.',
          constraint=PHONE, constraint_message=PHONE_MSG, appearance='numbers'),
        q('integer', 'hh_size', 'How many people usually live and eat in this household?',
          required='yes', constraint='. >= 1 and . <= 30', constraint_message='Between 1 and 30.'),
        q('end_group', 'respondent'),

        q('begin_repeat', 'person', 'Household member', repeat_count='${hh_size}'),
        q('calculate', 'person_no', calculation='position(..)'),
        q('text', 'p_name', 'Name of member ${person_no}', required='yes'),
        q('select_one relationship', 'p_relationship', 'Relationship to the head of household',
          required='yes'),
        q('select_one sex', 'p_sex', 'Sex', required='yes'),
        q('integer', 'p_age', 'Age in completed years', hint='Enter 0 for babies under one year.',
          required='yes', constraint='. >= 0 and . <= 120', constraint_message='Between 0 and 120.'),
        q('select_one marital', 'p_marital', 'Marital status', required='yes',
          relevant='${p_age} >= 12'),
        q('select_one yes_no', 'p_in_school', 'Currently attending school?', required='yes',
          relevant='${p_age} >= 3 and ${p_age} <= 30'),
        q('select_one education', 'p_education', 'Highest level of education completed',
          required='yes', relevant='${p_age} >= 5'),
        q('select_one activity', 'p_activity', 'Main activity in the last 7 days', required='yes',
          relevant='${p_age} >= 10'),
        q('select_one yes_no_dk', 'p_disability',
          'Does this person have difficulty seeing, hearing, walking or remembering?',
          required='yes'),
        q('select_one yes_no_dk', 'p_birth_registered', 'Is the birth registered?', required='yes',
          relevant='${p_age} < 5'),
        q('calculate', 'p_is_head', calculation="if(${p_relationship} = 'head', 1, 0)"),
        q('calculate', 'p_is_child', calculation='if(${p_age} < 15, 1, 0)'),
        q('end_repeat', 'person'),
        q('calculate', 'head_count', calculation='sum(${p_is_head})'),
        q('calculate', 'children_count', calculation='sum(${p_is_child})'),
        q('acknowledge', 'roster_check', 'Roster complete: ${hh_size} members, ${children_count} under 15.',
          required='yes', constraint='${head_count} = 1',
          constraint_message='The household must have exactly one head. Go back and correct the roster.'),

        q('begin_group', 'dwelling', 'Dwelling'),
        q('select_one dwelling', 'dwelling_type', 'Type of dwelling', required='yes'),
        q('select_one roof', 'roof', 'Main roof material', required='yes'),
        q('integer', 'rooms', 'Rooms used for sleeping', required='yes',
          constraint='. >= 1 and . <= 30', constraint_message='Between 1 and 30.'),
        q('select_one water', 'water_source', 'Main source of drinking water', required='yes'),
        q('select_one toilet', 'toilet', 'Type of toilet facility', required='yes'),
        q('select_multiple power', 'lighting', 'Sources of lighting', required='yes'),
        q('image', 'dwelling_photo', 'Photo of the dwelling (optional)', parameters='max-pixels=1024'),
        q('end_group', 'dwelling'),

        q('begin_group', 'events', 'Last 12 months'),
        q('integer', 'births_12m', 'Babies born alive in the household in the last 12 months',
          required='yes', constraint='. >= 0 and . <= 5', constraint_message='Between 0 and 5.'),
        q('integer', 'deaths_12m', 'Household members who died in the last 12 months',
          required='yes', constraint='. >= 0 and . <= 10', constraint_message='Between 0 and 10.'),
        q('end_group', 'events'),
        q('text', 'remarks', 'Enumerator remarks (optional)', appearance='multiline'),
        q('end_group', 'main'),
    ],
    'choices': {
        'yes_no': YES_NO, 'yes_no_dk': YES_NO_DK, 'sex': SEX, 'district': DISTRICTS,
        'relationship': RELATIONSHIP,
        'marital': [('never', 'Never married'), ('married', 'Married'),
                    ('cohabiting', 'Living together'), ('divorced', 'Divorced or separated'),
                    ('widowed', 'Widowed')],
        'education': [('none', 'None'), ('primary', 'Primary'), ('jss', 'Junior secondary'),
                      ('sss', 'Senior secondary'), ('tertiary', 'Tertiary'),
                      ('vocational', 'Vocational')],
        'activity': [('farming', 'Farming'), ('trading', 'Trading'), ('wage', 'Paid employment'),
                     ('mining', 'Mining'), ('fishing', 'Fishing'), ('student', 'Student'),
                     ('housework', 'Housework'), ('unemployed', 'Looking for work'),
                     ('none', 'Not working')],
        'dwelling': [('house', 'Separate house'), ('compound', 'Room in a compound'),
                     ('flat', 'Flat or apartment'), ('hut', 'Traditional hut'),
                     ('makeshift', 'Makeshift shelter')],
        'roof': [('zinc', 'Zinc sheets'), ('thatch', 'Thatch'), ('tiles', 'Tiles'),
                 ('concrete', 'Concrete'), ('tarpaulin', 'Tarpaulin'), ('other', 'Other')],
        'water': [('piped', 'Piped into dwelling or yard'), ('tap', 'Public tap'),
                  ('borehole', 'Borehole'), ('protected_well', 'Protected well'),
                  ('unprotected_well', 'Unprotected well'), ('river', 'River or stream'),
                  ('sachet', 'Sachet or bottled water')],
        'toilet': [('flush', 'Flush toilet'), ('vip', 'Ventilated pit latrine'),
                   ('pit_slab', 'Pit latrine with slab'), ('pit_open', 'Pit latrine without slab'),
                   ('none', 'No facility (bush or field)')],
        'power': [('grid', 'Electricity grid'), ('solar', 'Solar'), ('generator', 'Generator'),
                  ('torch', 'Battery torch'), ('kerosene', 'Kerosene lamp'), ('candle', 'Candle')],
    },
}

# A short back-check form for the census: same household, a few stable answers.
census_backcheck = {
    'file': 'fd_test_census_backcheck.xlsx',
    'settings': {'form_title': 'Test: Household census back-check',
                 'form_id': 'fd_test_census_backcheck',
                 'instance_name': "concat('BC ', ${household_id})"},
    'survey': start_meta() + consent(
        'Field Data TEST form. Census back-check: a short revisit to confirm a few answers. '
        'Takes about 5 minutes.') + [
        q('begin_group', 'main', '', relevant="${consent} = 'yes'"),
        q('text', 'enumerator', 'Back-checker name', required='yes'),
        *location('household_id', 'Household ID', '^HH-[0-9]{6}$',
                  'Use the format HH-000001.', 'The ID on the original census form.'),
        q('integer', 'hh_size', 'How many people usually live and eat in this household?',
          required='yes', constraint='. >= 1 and . <= 30', constraint_message='Between 1 and 30.'),
        q('select_one sex', 'head_sex', 'Sex of the head of household', required='yes'),
        q('integer', 'head_age', 'Age of the head of household', required='yes',
          constraint='. >= 12 and . <= 120', constraint_message='Between 12 and 120.'),
        q('select_one water', 'water_source', 'Main source of drinking water', required='yes'),
        q('select_one toilet', 'toilet', 'Type of toilet facility', required='yes'),
        q('select_one yes_no', 'was_visited', 'Did a census enumerator visit this household recently?',
          required='yes'),
        q('end_group', 'main'),
    ],
    'choices': {k: census['choices'][k] for k in ('yes_no', 'sex', 'district', 'water', 'toilet')},
}

# --------------------------------------------------------------------- health
TRACERS = [
    ('ors', 'ORS sachets'), ('zinc', 'Zinc tablets'), ('act', 'ACT antimalarial'),
    ('rdt', 'Malaria rapid test (RDT)'), ('amoxicillin', 'Amoxicillin'),
    ('oxytocin', 'Oxytocin'), ('iron_folate', 'Iron and folic acid'),
    ('penta', 'Pentavalent vaccine'),
]
health = {
    'file': 'fd_test_health.xlsx',
    'settings': {'form_title': 'Test: Health facility assessment',
                 'form_id': 'fd_test_health',
                 'instance_name': "concat(${facility_id}, ' - ', ${facility_name})"},
    'survey': start_meta() + consent(
        'Field Data TEST form. Health facility assessment. Ask for the facility in-charge. '
        'Takes about 25 minutes.') + [
        q('begin_group', 'main', '', relevant="${consent} = 'yes'"),
        q('text', 'enumerator', 'Assessor name', required='yes'),
        *location('facility_id', 'Facility ID', '^HF-[0-9]{4}$',
                  'Use the format HF-0001.', 'Format HF-0001, from the facility list.'),

        q('begin_group', 'facility', 'Facility'),
        q('text', 'facility_name', 'Facility name', required='yes'),
        q('select_one facility_type', 'facility_type', 'Facility type', required='yes'),
        q('select_one ownership', 'ownership', 'Managed by', required='yes'),
        q('text', 'incharge_name', 'Name of the in-charge', required='yes'),
        q('text', 'incharge_phone', 'In-charge phone number', required='yes',
          constraint=PHONE, constraint_message=PHONE_MSG, appearance='numbers'),
        q('select_one open_status', 'open_status', 'Is the facility open today?', required='yes'),
        q('image', 'facility_photo', 'Photo of the facility front', required='yes', parameters='max-pixels=1024'),
        q('end_group', 'facility'),

        q('begin_group', 'staff', 'Staff on the payroll', relevant="${open_status} != 'closed'"),
        q('integer', 'staff_doctors', 'Doctors', required='yes', default='0',
          constraint='. >= 0 and . <= 100', constraint_message='Between 0 and 100.'),
        q('integer', 'staff_cho', 'Community health officers', required='yes', default='0',
          constraint='. >= 0 and . <= 100', constraint_message='Between 0 and 100.'),
        q('integer', 'staff_nurses', 'Nurses', required='yes', default='0',
          constraint='. >= 0 and . <= 200', constraint_message='Between 0 and 200.'),
        q('integer', 'staff_midwives', 'Midwives', required='yes', default='0',
          constraint='. >= 0 and . <= 100', constraint_message='Between 0 and 100.'),
        q('integer', 'staff_present', 'How many of them are present today?', required='yes',
          constraint='. >= 0 and . <= ${staff_doctors} + ${staff_cho} + ${staff_nurses} + ${staff_midwives}',
          constraint_message='Cannot be more than the staff on the payroll.'),
        q('calculate', 'staff_total',
          calculation='${staff_doctors} + ${staff_cho} + ${staff_nurses} + ${staff_midwives}'),
        q('note', 'staff_note', 'Staff on payroll: ${staff_total}; present today: ${staff_present}'),
        q('end_group', 'staff'),

        q('begin_group', 'services_grp', 'Services and infrastructure',
          relevant="${open_status} != 'closed'"),
        q('select_multiple service', 'services', 'Which services does the facility offer?',
          required='yes'),
        q('select_one power', 'power', 'Main source of electricity', required='yes'),
        q('select_one water', 'water', 'Main source of water', required='yes'),
        q('select_one yes_no', 'toilet_working', 'Is there a working toilet for patients?',
          required='yes'),
        q('select_one yes_no', 'cold_chain', 'Is the vaccine fridge working today?', required='yes',
          relevant="selected(${services}, 'immunization')"),
        q('decimal', 'fridge_temp', 'Fridge temperature now (°C)', required='yes',
          relevant="${cold_chain} = 'yes'",
          constraint='. >= -30 and . <= 30', constraint_message='Between -30 and 30 °C.'),
        q('note', 'fridge_warn', 'Warning: vaccines should be kept between 2 and 8 °C.',
          relevant="${cold_chain} = 'yes' and (${fridge_temp} < 2 or ${fridge_temp} > 8)"),
        q('end_group', 'services_grp'),

        q('begin_group', 'stock', 'Tracer items in stock today', appearance='field-list',
          relevant="${open_status} != 'closed'"),
        q('select_one stock', 'stock_header', 'Item', appearance='label'),
        *[q('select_one stock', f'stock_{code}', label, required='yes', appearance='list-nolabel')
          for code, label in TRACERS],
        q('end_group', 'stock'),
        q('calculate', 'stockouts', calculation=' + '.join(
            f"if(${{stock_{code}}} = 'out', 1, 0)" for code, _ in TRACERS)),
        q('note', 'stockout_note', 'Tracer items out of stock: ${stockouts} of 8',
          relevant="${open_status} != 'closed'"),

        q('begin_group', 'last_month', 'Register totals for last month',
          relevant="${open_status} != 'closed'"),
        q('integer', 'opd_visits', 'Outpatient visits', required='yes',
          constraint='. >= 0 and . <= 20000', constraint_message='Between 0 and 20,000.'),
        q('integer', 'malaria_tested', 'People tested for malaria', required='yes',
          constraint='. >= 0 and . <= ${opd_visits}',
          constraint_message='Cannot be more than the outpatient visits.'),
        q('integer', 'malaria_positive', 'Malaria tests that were positive', required='yes',
          constraint='. >= 0 and . <= ${malaria_tested}',
          constraint_message='Cannot be more than the people tested.'),
        q('integer', 'anc_first', 'First antenatal care visits', required='yes',
          relevant="selected(${services}, 'anc')",
          constraint='. >= 0 and . <= 2000', constraint_message='Between 0 and 2,000.'),
        q('integer', 'deliveries', 'Deliveries at the facility', required='yes',
          relevant="selected(${services}, 'delivery')",
          constraint='. >= 0 and . <= 2000', constraint_message='Between 0 and 2,000.'),
        q('calculate', 'positivity',
          calculation='if(${malaria_tested} > 0, round(${malaria_positive} div ${malaria_tested} * 100, 1), 0)'),
        q('note', 'positivity_note', 'Malaria test positivity: ${positivity}%'),
        q('image', 'register_photo', 'Photo of the register page with these totals (optional)', parameters='max-pixels=1024'),
        q('end_group', 'last_month'),
        q('text', 'remarks', 'Assessor remarks (optional)', appearance='multiline'),
        q('end_group', 'main'),
    ],
    'choices': {
        'yes_no': YES_NO, 'district': DISTRICTS,
        'facility_type': [('mchp', 'Maternal and child health post (MCHP)'),
                          ('chp', 'Community health post (CHP)'),
                          ('chc', 'Community health centre (CHC)'),
                          ('hospital', 'Hospital'), ('clinic', 'Private clinic')],
        'ownership': [('government', 'Government'), ('mission', 'Mission or faith-based'),
                      ('ngo', 'NGO'), ('private', 'Private')],
        'open_status': [('open', 'Open'), ('partly', 'Open with reduced services'),
                        ('closed', 'Closed')],
        'service': [('opd', 'Outpatient care'), ('anc', 'Antenatal care'),
                    ('delivery', 'Deliveries'), ('immunization', 'Immunization'),
                    ('family_planning', 'Family planning'), ('malaria', 'Malaria testing and treatment'),
                    ('nutrition', 'Nutrition screening'), ('lab', 'Laboratory'),
                    ('inpatient', 'Inpatient beds')],
        'power': [('grid', 'Electricity grid'), ('solar', 'Solar'), ('generator', 'Generator'),
                  ('none', 'None')],
        'water': [('piped', 'Piped'), ('borehole', 'Borehole'), ('well', 'Well'),
                  ('rain', 'Rainwater tank'), ('none', 'None on site')],
        'stock': [('in', 'In stock'), ('low', 'Low'), ('out', 'Out of stock'),
                  ('never', 'Never stocked')],
    },
}

FORMS = [agriculture, census, census_backcheck, health]

FONT = Font(name='Arial', size=10)
HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='2F5D50')
GROUP_FILL = PatternFill('solid', fgColor='E8F0EC')


def sheet(wb, title, cols, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for row in rows:
        ws.append([row.get(c, '') for c in cols])
    for cell in ws[1]:
        cell.font, cell.fill = HEAD, FILL
    for r in ws.iter_rows(min_row=2):
        grouping = title == 'survey' and str(r[0].value).split(' ')[0] in (
            'begin_group', 'end_group', 'begin_repeat', 'end_repeat')
        for cell in r:
            cell.font = Font(name='Arial', size=10, bold=grouping)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if grouping:
                cell.fill = GROUP_FILL
    for i, col in enumerate(cols):
        ws.column_dimensions[ws.cell(row=1, column=i + 1).column_letter].width = widths.get(col, 18)
    ws.freeze_panes = 'C2'
    return ws


def build(form):
    wb = Workbook()
    wb.remove(wb.active)
    used = [c for c in SURVEY_COLS if any(r.get(c) for r in form['survey'])]
    sheet(wb, 'survey', used, form['survey'],
          {'type': 26, 'name': 20, 'label': 48, 'hint': 36, 'relevant': 32, 'constraint': 36,
           'constraint_message': 36, 'calculation': 40, 'choice_filter': 28})
    choices = [{'list_name': ln, 'name': n, 'label': lab}
               for ln, items in form['choices'].items() for n, lab in items]
    sheet(wb, 'choices', ['list_name', 'name', 'label'], choices,
          {'list_name': 18, 'name': 20, 'label': 40})
    settings = dict(form['settings'], version=VERSION)
    sheet(wb, 'settings', list(settings), [settings], {'form_title': 40, 'instance_name': 48})
    wb.save(HERE / form['file'])


if __name__ == '__main__':
    for form in FORMS:
        build(form)
        print('wrote', form['file'])
