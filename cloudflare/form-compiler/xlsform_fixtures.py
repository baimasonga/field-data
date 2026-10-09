"""Real XLSForm workbooks shared by the contract tests and the runtime check.

They are built in code so no binary fixture is committed, and they cover the
constructs the compiler must keep working: skip logic, constraints, repeat
groups and GPS capture. Nothing here mocks pyxform.
"""

import io

from openpyxl import Workbook

SURVEY_HEADER = ["type", "name", "label", "relevant", "constraint", "constraint_message", "required"]


def xlsx_bytes(survey, choices=None, settings=None):
    """Serialise sheets to .xlsx bytes. Each argument is a list of rows, header first."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "survey"
    for row in survey:
        sheet.append(row)
    if choices:
        sheet = workbook.create_sheet("choices")
        for row in choices:
            sheet.append(row)
    if settings:
        sheet = workbook.create_sheet("settings")
        for row in settings:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def household_workbook(form_id="household"):
    """Skip logic, a numeric constraint, a repeat group and a GPS point."""
    return xlsx_bytes(
        [
            SURVEY_HEADER,
            ["text", "respondent", "Respondent name", "", "", "", "yes"],
            ["integer", "age", "Age", "", ". >= 0 and . <= 120", "Age must be between 0 and 120", ""],
            ["select_one yn", "has_children", "Does the household have children?", "", "", "", ""],
            ["begin repeat", "children", "Children", '${has_children} = "yes"', "", "", ""],
            ["text", "child_name", "Child name", "", "", "", "yes"],
            ["end repeat", "", "", "", "", "", ""],
            ["geopoint", "location", "Household location", "", "", "", ""],
        ],
        choices=[["list_name", "name", "label"], ["yn", "yes", "Yes"], ["yn", "no", "No"]],
        settings=[["form_title", "form_id"], ["Household survey", form_id]],
    )


def undefined_reference_workbook():
    """An XLSForm whose author referenced a question that does not exist."""
    return xlsx_bytes(
        [
            SURVEY_HEADER,
            ["text", "respondent", "Respondent name", "${no_such_question} = 'x'", "", "", ""],
        ],
        settings=[["form_title", "form_id"], ["Broken", "broken"]],
    )


def unknown_type_workbook():
    """An XLSForm using a question type that does not exist."""
    return xlsx_bytes(
        [SURVEY_HEADER, ["not_a_real_type", "q", "Question", "", "", "", ""]],
        settings=[["form_title", "form_id"], ["Broken type", "broken_type"]],
    )
