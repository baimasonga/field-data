# Builds the XLSForms behind planted.xml and corrected.xml; compile each with
# pyxform's xls2xform (--skip_validate, since planted.xml has a deliberate cycle).
import openpyxl, sys
def book(survey, choices, settings, path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'survey'
    for row in survey: ws.append(row)
    c = wb.create_sheet('choices')
    for row in choices: c.append(row)
    s = wb.create_sheet('settings')
    for row in settings: s.append(row)
    wb.save(path)

H = ['type','name','label::English (en)','label::Krio (kri)','relevant','constraint','calculation','required','read_only','choice_filter','default']
bad = [H,
 ['select_one district','district','District','Distrikt','','','','yes','','',''],
 ['select_one chiefdom','chiefdom','Chiefdom','Chiefdom','','','','','',"district=${district}",''],
 ['integer','age','Age','Ej','','. > 18 and . < 16','','','','',''],
 ['integer','hh_size','Household size','','','','','','','',''],
 ['text','bo_only','Only in Bo','Na Bo nomo',"${district} = 'bo'",'','','','','',''],
 ['select_multiple assets','assets','Assets','Tin dem','','','','','','',''],
 ['text','tv_brand','TV brand','TV brand',"selected(${assets}, 'television')",'','','','','',''],
 ['text','early','Asked early','Asked early',"${late} != ''",'','','','','',''],
 ['text','late','Asked late','Asked late','','','','','','',''],
 ['calculate','a','','','','','${b} + 1','','','',''],
 ['calculate','b','','','','','${a} + 1','','','',''],
 ['text','locked','Locked','Locked','','','','yes','yes','',''],
 ['begin_repeat','member','Member','Memba','','','','','','',''],
 ['text','m_name','Name','Nem','','','','','','',''],
 ['integer','m_age','Age','Ej',"${m_name} != ''",'. >= 0 and . <= 120','','','','',''],
 ['end_repeat','','','','','','','','','',''],
 ['begin_group','never','Never','Never','false()','','','','','',''],
 ['text','inside','Inside','Inside','','','','','','',''],
 ['end_group','','','','','','','','','',''],
 ['text','fancy','Fancy','Fancy','string-length(${late}) > 3','','','','','',''],
]
CH = ['list_name','name','label::English (en)','label::Krio (kri)','district']
choices = [CH,
 ['district','Bo','Bo','Bo',''],['district','Kenema','Kenema','Kenema',''],
 ['chiefdom','badjia','Badjia','Badjia','Bo'],['chiefdom','nongowa','Nongowa','Nongowa','Kenema'],
 ['assets','tv','TV','TV',''],['assets','radio','Radio','Radio',''],['assets','phone','Radio','Fon',''],
]
book(bad, choices, [['form_id','version'],['doctor_bad','1']], 'bad.xlsx')
good = [r[:] for r in bad]
# corrected version
fix = {'age': (5, '. >= 0 and . <= 120'), 'bo_only': (4, "${district} = 'Bo'"), 'tv_brand': (4, "selected(${assets}, 'tv')")}
for r in good:
    if r[1] in fix: r[fix[r[1]][0]] = fix[r[1]][1]
    if r[1] == 'hh_size': r[3] = 'Haus sayz'; r[5] = '. >= 1 and . <= 50'
    if r[1] == 'early': r[4] = ''
    if r[1] == 'b': r[6] = '1'
    if r[1] == 'locked': r[8] = ''
    if r[1] == 'never': r[4] = ''
good_choices = [r[:] for r in choices]
good_choices[-1][2] = 'Phone'
book(good, good_choices, [['form_id','version'],['doctor_good','1']], 'good.xlsx')
